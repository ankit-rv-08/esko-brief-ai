"""Convert supported source files into citeable SourceElements."""
from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any, Optional

import pymupdf
from docx import Document
from openpyxl import load_workbook

from src.models import SourceElement, SourceType
from src.prompts import VISION_EXTRACTION_SYSTEM, VISION_EXTRACTION_USER

logger = logging.getLogger(__name__)
_FRAGMENTED_THRESHOLD = 0.35


def _looks_fragmented(text: str) -> bool:
    tokens = text.split()
    if len(tokens) < 20:
        return False
    return sum(len(token) == 1 for token in tokens) / len(tokens) > _FRAGMENTED_THRESHOLD


def _dehyphenate(text: str) -> str:
    return re.sub(r"(\w)-\s*\n\s*(\w)", r"\1\2", text)


def ingest_pdf(
    path: Path,
    *,
    render_pages: bool = True,
    describe_with_vision: bool = False,
    vision_client: Optional[Any] = None,
) -> list[SourceElement]:
    """Extract one element per PDF page, optionally rendering pages for vision."""
    elements: list[SourceElement] = []
    image_dir = path.parent / "_rendered" / path.stem
    if render_pages:
        image_dir.mkdir(parents=True, exist_ok=True)

    with pymupdf.open(str(path)) as document:
        for page_index, page in enumerate(document):
            page_number = page_index + 1
            text = _dehyphenate(page.get_text()).strip()
            image_path: Optional[str] = None
            if render_pages:
                image_path = str(image_dir / f"page_{page_number}.png")
                page.get_pixmap(dpi=150).save(image_path)
            if describe_with_vision and vision_client and _looks_fragmented(text):
                try:
                    text = vision_client.describe(
                        image_path=image_path,
                        system=VISION_EXTRACTION_SYSTEM,
                        user=VISION_EXTRACTION_USER,
                    )
                except Exception as exc:  # noqa: BLE001
                    logger.warning("Vision fallback failed for page %d: %s", page_number, exc)
            if not text:
                text = "(no extractable text - see rendered page image)"
            elements.append(SourceElement(
                source_file=path.name, source_type=SourceType.PDF,
                page=page_number, content=text, image_path=image_path,
            ))
    return elements


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def ingest_xlsx(path: Path) -> list[SourceElement]:
    """Parse every non-empty workbook row into one source element."""
    elements: list[SourceElement] = []
    workbook = load_workbook(str(path), data_only=True)
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [_cell_to_str(value) for value in row]
            if any(cells):
                elements.append(SourceElement(
                    source_file=path.name, source_type=SourceType.XLSX,
                    sheet=sheet_name, row=row_number,
                    content=" | ".join(cell for cell in cells if cell),
                ))
    return elements


def ingest_docx(path: Path) -> list[SourceElement]:
    """Parse meaningful DOCX paragraphs into source elements."""
    elements: list[SourceElement] = []
    for paragraph_number, paragraph in enumerate(Document(str(path)).paragraphs, start=1):
        text = paragraph.text.strip()
        if len(text) >= 20:
            elements.append(SourceElement(
                source_file=path.name, source_type=SourceType.DOCX,
                row=paragraph_number, content=text,
            ))
    return elements


def ingest_image(path: Path, *, vision_client: Optional[Any] = None) -> list[SourceElement]:
    """Create one image element, describing it when a vision client is supplied."""
    content = "(image - no description available)"
    if vision_client:
        try:
            content = vision_client.describe(
                image_path=str(path), system=VISION_EXTRACTION_SYSTEM, user=VISION_EXTRACTION_USER
            )
        except Exception as exc:  # noqa: BLE001
            logger.warning("Vision description failed for %s: %s", path.name, exc)
    return [SourceElement(
        source_file=path.name, source_type=SourceType.IMAGE,
        content=content, image_path=str(path),
    )]


def ingest_directory(
    directory: Path,
    *,
    vision_client: Optional[Any] = None,
    describe_pdfs_with_vision: bool = False,
) -> list[SourceElement]:
    """Ingest supported files in deterministic filename order."""
    elements: list[SourceElement] = []
    handlers = {".pdf": ingest_pdf, ".xlsx": ingest_xlsx, ".docx": ingest_docx}
    image_extensions = {".png", ".jpg", ".jpeg"}
    for path in sorted(directory.iterdir()):
        if path.is_dir():
            continue
        try:
            suffix = path.suffix.lower()
            if suffix in handlers:
                if suffix == ".pdf":
                    elements.extend(ingest_pdf(
                        path, render_pages=True,
                        describe_with_vision=describe_pdfs_with_vision,
                        vision_client=vision_client,
                    ))
                else:
                    elements.extend(handlers[suffix](path))
            elif suffix in image_extensions:
                elements.extend(ingest_image(path, vision_client=vision_client))
        except Exception:  # noqa: BLE001
            logger.exception("Failed to ingest %s", path)
    return elements